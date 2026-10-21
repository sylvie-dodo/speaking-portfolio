"""명렬 엑셀(.xlsx) 또는 CSV(열 순번·학번·학년·이름)로 students를 적재하고 settings 1행을 만든다.

학년이 빈 행 하나를 교사로 지정한다. 학생은 학번과 이름으로 입장하며, 순번은 내부 번호로 쓴다.

사용: python scripts/seed_students.py "<명렬 파일>"          (검증만)
      python scripts/seed_students.py "<명렬 파일>" --apply  (서버에 적재)
"""
import csv
import sys
from pathlib import Path

CLASS = 1
COLUMNS = ("순번", "학번", "학년", "이름")


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        values = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        header, body = values[0], values[1:]
        return [dict(zip(header, r)) for r in body if any(r)]
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [{k: _cell(v) for k, v in r.items()} for r in csv.DictReader(f) if any((v or "").strip() for v in r.values())]


def read_roster(path: Path) -> list[dict]:
    raw = read_rows(path)
    if not raw or any(c not in raw[0] for c in COLUMNS):
        raise SystemExit(f"열 이름은 {', '.join(COLUMNS)}이어야 합니다.")
    rows = []
    for r in raw:
        if not r["순번"].isdigit() or not r["학번"] or not r["이름"]:
            raise SystemExit(f"순번·학번·이름이 비었거나 잘못된 행: 순번 {r['순번']!r}")
        grade = r["학년"]
        if grade and not grade.isdigit():
            raise SystemExit(f"순번 {r['순번']}: 학년은 숫자여야 합니다.")
        rows.append({"class": CLASS, "number": int(r["순번"]), "student_id": r["학번"].upper(), "name": r["이름"],
                     "grade": int(grade) if grade else None, "is_teacher": not grade})
    for key in ("number", "student_id"):
        values = [r[key] for r in rows]
        if len(set(values)) != len(values):
            raise SystemExit(f"중복된 {'순번' if key == 'number' else '학번'}이 있습니다.")
    if sum(r["is_teacher"] for r in rows) != 1:
        raise SystemExit("교사 행(학년이 빈 행)이 정확히 하나여야 합니다.")
    return rows


def main():
    args = [a for a in sys.argv[1:] if a != "--apply"]
    if len(args) != 1:
        raise SystemExit(__doc__)
    rows = read_roster(Path(args[0]))
    students = [r for r in rows if not r["is_teacher"]]
    grades = sorted({r["grade"] for r in students})
    print(f"학생 {len(students)}명(학년 {', '.join(map(str, grades))}), 교사 1명, 순번 {rows[0]['number']}~{rows[-1]['number']}")
    if "--apply" not in sys.argv:
        print("검증만 했습니다. 서버에 올리려면 --apply를 붙이세요.")
        return
    from common import service_client
    client = service_client()
    client.table("students").upsert(rows).execute()
    client.table("settings").upsert({"class": CLASS, "exam_open": False}).execute()
    print("서버 적재 완료(settings 1행 포함)")


if __name__ == "__main__":
    main()
